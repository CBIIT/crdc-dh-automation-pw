"""PBACPermissions.xlsx is the expected-permissions driver (Katalon Data File CRDC/Pbac/Permissions).

write_expected copies one role column into OutputFiles/PBAC_Defaults_Results.xlsx.
compare_actual fills Actual / Actual Name / Result against UI checkbox rows in sheet order.
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType

from crdc_pw.locators.loader import repo_root

log = logging.getLogger("crdc_pw.pbac_sheet")

INPUT_RELATIVE = Path("InputFiles/CRDC/PBACPermissions.xlsx")
OUTPUT_RELATIVE = Path("OutputFiles/PBAC_Defaults_Results.xlsx")
SHEET_NAME = "Permissions"
INTERNAL_COL = "Permission Internal Name"
DISPLAY_COL = "Name"

HEADERS = ("Permission", "Expected", "Actual", "Name", "Actual Name", "Result")

PASS_FILL = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER_FONT = Font(bold=True)


def input_path(root: Path | None = None) -> Path:
    return (root or repo_root()) / INPUT_RELATIVE


def output_path(root: Path | None = None) -> Path:
    return (root or repo_root()) / OUTPUT_RELATIVE


def load_expected_rows(role: str, root: Path | None = None) -> list[tuple[str, str, str]]:
    """Rows for one role column: (internal_name, expected_status, display_name).

    Skips rows with a blank Permission Internal Name, same as Katalon writePbacExpectedPermissionsForRole.
    """
    path = input_path(root)
    wb = load_workbook(path)
    ws = wb[SHEET_NAME]
    headers = [_cell_str(ws.cell(1, col)) for col in range(1, ws.max_column + 1)]
    try:
        role_idx = headers.index(role) + 1
    except ValueError as exc:
        raise KeyError(f"Column {role!r} not found in {path.name} sheet {SHEET_NAME}") from exc
    internal_idx = headers.index(INTERNAL_COL) + 1
    display_idx = headers.index(DISPLAY_COL) + 1

    rows: list[tuple[str, str, str]] = []
    for excel_row in range(2, ws.max_row + 1):
        internal = _cell_str(ws.cell(excel_row, internal_idx))
        if not internal:
            continue
        expected = _cell_str(ws.cell(excel_row, role_idx))
        display = _cell_str(ws.cell(excel_row, display_idx))
        rows.append((internal, expected, display))
    return rows


def write_expected(role: str, root: Path | None = None) -> Path:
    """Write the role's expected column into OutputFiles/PBAC_Defaults_Results.xlsx."""
    expected_rows = load_expected_rows(role, root)
    out = output_path(root)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = _open_or_create(out)
    if role in wb.sheetnames:
        del wb[role]
    ws = wb.create_sheet(role)
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(1, col, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for i, (internal, expected, display) in enumerate(expected_rows, start=2):
        ws.cell(i, 1, internal)
        ws.cell(i, 2, expected)
        ws.cell(i, 4, display)
    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 22
    _save(wb, out)
    log.info("Expected permissions for %s written to %s", role, out)
    return out


def compare_actual(
    role: str,
    actual_rows: list[tuple[str, str]],
    root: Path | None = None,
) -> bool:
    """Fill Actual / Actual Name / Result from UI rows in sheet order. True if every row Pass.

    actual_rows is (display_label, status) in checkbox order, same as Katalon i-1 vs sheet row i.
    """
    out = output_path(root)
    wb = load_workbook(out)
    if role not in wb.sheetnames:
        raise KeyError(f"Sheet {role!r} not found in {out}")
    ws = wb[role]
    overall = True
    last = ws.max_row
    for excel_row in range(2, last + 1):
        expected = _cell_str(ws.cell(excel_row, 2))
        expected_name = _cell_str(ws.cell(excel_row, 4))
        ui_index = excel_row - 2
        if ui_index >= len(actual_rows):
            actual_status = "MISSING_UI"
            ui_label = "MISSING_UI"
            passed = False
        else:
            ui_label, actual_status = actual_rows[ui_index]
            if not ui_label:
                ui_label = "MISSING_LABEL_UI"
            state_matches = bool(expected) and expected.lower() == (actual_status or "").lower()
            name_matches = bool(expected_name) and bool(ui_label) and expected_name.lower() == ui_label.lower()
            passed = state_matches and name_matches
        ws.cell(excel_row, 3, actual_status)
        ws.cell(excel_row, 5, ui_label)
        result_cell = ws.cell(excel_row, 6, "Pass" if passed else "Fail")
        result_cell.fill = PASS_FILL if passed else FAIL_FILL
        if not passed:
            overall = False
        log.info(
            "Row %s - %s | ExpectedState: %s | ActualState: %s | ExpectedName: %s | ActualName: %s | Result: %s",
            excel_row,
            expected_name,
            expected,
            actual_status,
            expected_name,
            ui_label,
            "Pass" if passed else "Fail",
        )
    _save(wb, out)
    if overall:
        log.info("Actual permissions for %r recorded and compared.", role)
    else:
        log.error("There is a FAILURE -- verify in output Excel: %s", out)
    return overall


def _open_or_create(path: Path) -> WorkbookType:
    if path.is_file():
        return load_workbook(path)
    return Workbook()


def _save(wb: WorkbookType, path: Path) -> None:
    extra = [name for name in wb.sheetnames if name.startswith("Sheet")]
    if extra and len(wb.sheetnames) > 1:
        for name in extra:
            del wb[name]
    wb.save(path)
    wb.close()


def _cell_str(cell: object) -> str:
    value = getattr(cell, "value", None)
    if value is None:
        return ""
    return str(value).strip()
