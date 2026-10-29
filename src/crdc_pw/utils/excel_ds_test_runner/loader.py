"""Load Meta and Steps from CRDC Excel test case workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from crdc_pw.config import ROLES
from crdc_pw.locators.loader import repo_root

META_SHEET = "Meta"
STEPS_SHEET = "Steps"

_COL_INPUT = 0
_COL_ACTION = 1
_COL_PARAMS = 2
_COL_MODEL = 3

_STEPS_HEADER = ("input excel", "action", "params", "model")


@dataclass(frozen=True)
class StepRow:
    """One executable row on the Steps sheet (row 3+)."""

    action: str
    params: str
    model: str
    excel_row: int


class ExcelLoadError(ValueError):
    """Excel test case file failed validation."""

    def __init__(self, message: str, *, path: Path | None = None, row: int | None = None) -> None:
        loc = ""
        if path is not None:
            loc = str(path)
            if row is not None:
                loc += f" row {row}"
        super().__init__(f"{message}" + (f" ({loc})" if loc else ""))


ScenarioLoadError = ExcelLoadError


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_workbook_rows(path: Path, sheet: str) -> list[tuple[object, ...]]:
    if not path.is_file():
        raise ExcelLoadError("Workbook not found", path=path)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ExcelLoadError(f"Missing sheet '{sheet}'", path=path)
        return list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()


def read_meta(path: Path) -> dict[str, str]:
    """Parse Meta sheet as key/value pairs (column A = key, B = value)."""
    rows = _read_workbook_rows(path, META_SHEET)
    meta: dict[str, str] = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        key = _cell_str(row[0]).lower()
        if not key or key == "key":
            continue
        val = _cell_str(row[1]) if len(row) > 1 else ""
        meta[key] = val
    return meta


def read_meta_role(path: Path) -> str:
    """Return hub role from Meta sheet; validates against ROLES."""
    meta = read_meta(path)
    role = meta.get("role", "")
    if not role:
        raise ExcelLoadError("Meta sheet must include role", path=path)
    if role not in ROLES:
        raise ExcelLoadError(
            f"Invalid Meta role '{role}'; expected one of: {', '.join(ROLES)}",
            path=path,
        )
    return role


def _input_path_from_config_row(cfg: tuple[object, ...]) -> str:
    if not cfg:
        return ""
    if len(cfg) > _COL_INPUT and cfg[_COL_INPUT]:
        return _cell_str(cfg[_COL_INPUT])
    return ""


def load_steps(path: Path) -> tuple[str, list[StepRow]]:
    """Return (input_root from row 2, executable step rows from row 3+)."""
    rows = _read_workbook_rows(path, STEPS_SHEET)
    if len(rows) < 2:
        raise ExcelLoadError("Steps sheet must have header and input row", path=path)

    header = tuple(_cell_str(c).lower() for c in rows[0])
    for i, name in enumerate(_STEPS_HEADER):
        if i >= len(header) or header[i] != name:
            raise ExcelLoadError(
                "Steps row 1 must be: Input Excel | Action | Params | Model",
                path=path,
            )

    input_root = _input_path_from_config_row(rows[1])
    if input_root:
        resolved = repo_root() / input_root.replace("\\", "/")
        if not resolved.is_dir():
            raise ExcelLoadError(f"Input path is not a directory: {input_root}", path=path)

    steps: list[StepRow] = []
    for excel_row, row in enumerate(rows[2:], start=3):
        if not row or len(row) <= _COL_ACTION or not row[_COL_ACTION]:
            continue
        steps.append(
            StepRow(
                action=_cell_str(row[_COL_ACTION]),
                params=_cell_str(row[_COL_PARAMS]) if len(row) > _COL_PARAMS else "",
                model=_cell_str(row[_COL_MODEL]) if len(row) > _COL_MODEL else "",
                excel_row=excel_row,
            )
        )
    if not steps:
        raise ExcelLoadError("Steps sheet has no action rows", path=path)
    return input_root, steps
