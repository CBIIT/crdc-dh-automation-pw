"""No-browser checks that PBACPermissions.xlsx drives expected rows and Pass/Fail compare.

Does not open Hub. Copies the committed xlsx into a temp tree so OutputFiles stay off the repo.

Flow (this file):
  1. Fedlead column loads from the committed workbook
  2. Matching UI rows vs that column are all Pass
  3. One wrong checkbox state is Fail
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from crdc_pw.locators.loader import repo_root
from crdc_pw.utils.pbac_sheet import (
    INPUT_RELATIVE,
    compare_actual,
    load_expected_rows,
    output_path,
    write_expected,
)


def _temp_input_tree(tmp_path: Path) -> Path:
    dest = tmp_path / INPUT_RELATIVE
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(repo_root() / INPUT_RELATIVE, dest)
    return tmp_path


@pytest.mark.unit
def test_fedlead_column_loads_from_committed_xlsx() -> None:
    """Sheet Permissions has a Fedlead column; first row is Submission Request View, fixed_checked."""
    rows = load_expected_rows("Fedlead")
    assert len(rows) >= 40
    internal, expected, display = rows[0]
    assert internal == "submission_request:view"
    assert expected == "fixed_checked"
    assert display == "View"


@pytest.mark.unit
def test_compare_matching_rows_all_pass(tmp_path: Path) -> None:
    """UI labels and statuses copied from the Fedlead column must all Pass."""
    root = _temp_input_tree(tmp_path)
    write_expected("Fedlead", root=root)
    expected = load_expected_rows("Fedlead", root=root)
    actual = [(display, status) for _internal, status, display in expected]
    assert compare_actual("Fedlead", actual, root=root)
    wb = load_workbook(output_path(root))
    ws = wb["Fedlead"]
    results = [ws.cell(row, 6).value for row in range(2, ws.max_row + 1)]
    assert results
    assert all(value == "Pass" for value in results)


@pytest.mark.unit
def test_compare_mismatched_state_fails(tmp_path: Path) -> None:
    """One checkbox with the wrong status marks that row Fail and the compare returns False."""
    root = _temp_input_tree(tmp_path)
    write_expected("Fedlead", root=root)
    expected = load_expected_rows("Fedlead", root=root)
    actual = [(display, status) for _internal, status, display in expected]
    label, _status = actual[0]
    actual[0] = (label, "unchecked")
    assert compare_actual("Fedlead", actual, root=root) is False
    wb = load_workbook(output_path(root))
    ws = wb["Fedlead"]
    assert ws.cell(2, 6).value == "Fail"
    assert ws.cell(2, 3).value == "unchecked"
