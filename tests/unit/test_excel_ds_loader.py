"""Unit tests for Excel DS Settings/Steps loader and step runner (no browser)."""

from __future__ import annotations

from pathlib import Path

import pytest

import shutil

from crdc_pw.config import Settings, load_settings, profile_paths
from crdc_pw.utils.excel_ds_test_runner.config_helpers import excel_ds_cases_glob, excel_ds_enabled
from crdc_pw.utils.excel_ds_test_runner.loader import ExcelLoadError, load_steps, read_settings_role
from crdc_pw.utils.excel_ds_test_runner.step_runner import KNOWN_ACTIONS, run_testcase_excel
from crdc_pw.locators.loader import repo_root

CANONICAL_VR_SEQUENCE = (
    "open_data_submissions",
    "create_data_submission",
    "upload_metadata_folder",
    "click_validate",
    "verify_validation_results",
)


def _norm_action(action: str) -> str:
    return action.strip().lower().replace(" ", "_").replace("-", "_")


def _discover_committed_excel_cases() -> list[Path]:
    root = repo_root()
    paths = sorted(
        p
        for p in root.glob("tests/DataValidation/**/TC*.xlsx")
        if p.is_file()
        and not p.name.startswith("~$")
        and not p.name.startswith("TC_Template")
    )
    return paths


def _validate_case_workbook(path: Path) -> None:
    """Offline checks for committed TC*.xlsx (same Params/Model rules as step_runner)."""
    read_settings_role(path)
    input_root, steps = load_steps(path)
    if not steps:
        pytest.fail(f"{path}: Steps sheet has no action rows")

    normalized = [_norm_action(s.action) for s in steps]
    if normalized != list(CANONICAL_VR_SEQUENCE):
        pytest.fail(
            f"{path}: expected VR step sequence {CANONICAL_VR_SEQUENCE}, got {tuple(normalized)}"
        )

    upload_folder: str | None = None
    verify_case_id: str | None = None
    verify_model: str | None = None
    create_commons: str | None = None

    for step in steps:
        action = _norm_action(step.action)
        params = step.params.strip()
        model = step.model.strip()
        loc = f"{path} row {step.excel_row} action {step.action!r}"

        if action not in KNOWN_ACTIONS:
            pytest.fail(f"{loc}: unknown action; use one of: {', '.join(KNOWN_ACTIONS)}")

        if action in ("open_data_submissions", "click_validate"):
            if params or model:
                pytest.fail(f"{loc}: Params and Model must be empty")
        elif action == "create_data_submission":
            if model:
                pytest.fail(f"{loc}: Model must be empty (use Params with four pipe-separated values)")
            parts = [p.strip() for p in params.split("|")]
            if len(parts) != 4:
                pytest.fail(f"{loc}: Params need four pipe-separated values, got {len(parts)}")
            create_commons = parts[2]
        elif action == "upload_metadata_folder":
            if model:
                pytest.fail(f"{loc}: Model must be empty")
            if not params:
                pytest.fail(f"{loc}: Params must be the MetadataData folder name (e.g. TC08)")
            upload_folder = params
        elif action == "verify_validation_results":
            if not params or not model:
                pytest.fail(f"{loc}: Params = case folder id, Model = commons (GC, ICDC, CTDC)")
            verify_case_id = params
            verify_model = model

    assert upload_folder is not None and verify_case_id is not None and verify_model is not None
    if upload_folder != verify_case_id:
        pytest.fail(
            f"{path}: upload_metadata_folder Params {upload_folder!r} "
            f"must match verify_validation_results Params {verify_case_id!r}"
        )
    expected_suffix = f"MetadataData/{verify_case_id}".replace("\\", "/")
    if not input_root.replace("\\", "/").endswith(expected_suffix):
        pytest.fail(
            f"{path}: row 2 input root {input_root!r} must end with {expected_suffix!r}"
        )
    if create_commons and create_commons != verify_model:
        pytest.fail(
            f"{path}: create_data_submission commons {create_commons!r} "
            f"must match verify_validation_results Model {verify_model!r}"
        )


def _write_minimal_excel_case(
    out_path: Path,
    *,
    role: str,
    input_root: str,
    steps: list[tuple[str, str, str]],
) -> None:
    """Minimal Settings + Steps workbook for loader/runner unit tests only."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Steps"
    ws.append(["Input Excel", "Action", "Params", "Model"])
    ws.append([input_root, None, None, None])
    for action, params, model in steps:
        ws.append([None, action, params, model])
    settings_ws = wb.create_sheet("Settings")
    settings_ws.append(["key", "value"])
    settings_ws.append(["role", role])
    settings_ws.append(["enabled", "Y"])
    wb.save(out_path)


@pytest.mark.unit
@pytest.mark.parametrize(
    "case_path",
    _discover_committed_excel_cases(),
    ids=lambda p: p.stem,
)
def test_committed_excel_case_workbook(case_path: Path) -> None:
    assert case_path.is_file(), case_path
    _validate_case_workbook(case_path)


@pytest.mark.unit
def test_missing_steps_sheet_raises(tmp_path: Path) -> None:
    path = tmp_path / "TC99_bad.xlsx"
    _write_minimal_excel_case(
        path,
        role="Submitter",
        input_root="InputFiles/CRDC/MetadataData/TC08",
        steps=[("open_data_submissions", "", "")],
    )
    import openpyxl

    wb = openpyxl.load_workbook(path)
    del wb["Steps"]
    wb.save(path)
    with pytest.raises(ExcelLoadError, match="Steps"):
        load_steps(path)


@pytest.mark.unit
def test_invalid_role_in_settings_raises(tmp_path: Path) -> None:
    path = tmp_path / "TC99_role.xlsx"
    _write_minimal_excel_case(
        path,
        role="NotARealRole",
        input_root="InputFiles/CRDC/MetadataData/TC08",
        steps=[("open_data_submissions", "", "")],
    )
    with pytest.raises(ExcelLoadError, match="Invalid Settings role"):
        read_settings_role(path)


@pytest.mark.unit
def test_run_testcase_excel_dispatches_steps(tmp_path: Path) -> None:
    path = tmp_path / "TC99_run.xlsx"
    _write_minimal_excel_case(
        path,
        role="Submitter",
        input_root="InputFiles/CRDC/MetadataData/TC08",
        steps=[
            ("open_data_submissions", "", ""),
            (
                "create_data_submission",
                "New|metadata only|GC|autoQA - AutomationQA_Study",
                "",
            ),
            ("upload_metadata_folder", "TC08", ""),
            ("click_validate", "", ""),
            ("verify_validation_results", "TC08", "GC"),
        ],
    )
    calls: list[tuple[str, ...]] = []

    class StubList:
        def open(self) -> None:
            calls.append(("open",))

        def create_data_submission(self, a: str, b: str, c: str, d: str) -> None:
            calls.append(("create", a, b, c, d))

    class StubDetail:
        def upload_folder(self, folder: str) -> None:
            calls.append(("upload", folder))

        def click_validate(self) -> None:
            calls.append(("validate",))

        def verify_validation_results_vs_expected_for_tc(self, tc: str, *, model: str) -> None:
            calls.append(("verify", tc, model))

    class StubHub:
        page = object()
        locators = object()

    import crdc_pw.utils.excel_ds_test_runner.step_runner as sr

    orig_list = sr.DataSubmissionsListPage
    orig_detail = sr.DataSubmissionDetailPage
    list_inst = StubList()
    detail_inst = StubDetail()
    sr.DataSubmissionsListPage = lambda *a, **k: list_inst  # type: ignore[misc, assignment]
    sr.DataSubmissionDetailPage = lambda *a, **k: detail_inst  # type: ignore[misc, assignment]
    try:
        run_testcase_excel(StubHub(), path)  # type: ignore[arg-type]
    finally:
        sr.DataSubmissionsListPage = orig_list
        sr.DataSubmissionDetailPage = orig_detail

    assert calls[0] == ("open",)
    assert calls[1] == ("create", "New", "metadata only", "GC", "autoQA - AutomationQA_Study")
    assert ("upload", "TC08") in calls
    assert ("validate",) in calls
    assert calls[-1] == ("verify", "TC08", "GC")


@pytest.mark.unit
def test_unknown_action_raises(tmp_path: Path) -> None:
    path = tmp_path / "TC99_unknown.xlsx"
    _write_minimal_excel_case(
        path,
        role="Submitter",
        input_root="InputFiles/CRDC/MetadataData/TC08",
        steps=[("not_a_real_action", "", "")],
    )

    class StubHub:
        page = object()
        locators = object()

    with pytest.raises(ExcelLoadError, match="Unknown step action"):
        run_testcase_excel(StubHub(), path)  # type: ignore[arg-type]


@pytest.mark.unit
def test_excel_ds_enabled_from_example_yaml(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("CRDC_EXCEL_DS_VALIDATION", raising=False)
    _, example = profile_paths(repo_root(), "qa")
    cfg_dir = tmp_path / "config"
    cfg_dir.mkdir()
    shutil.copy(example, cfg_dir / "qa.yaml.example")
    settings = load_settings(profile="qa", root=tmp_path)
    assert excel_ds_enabled(settings) is True
    assert "DataValidation" in excel_ds_cases_glob(settings)


@pytest.mark.unit
def test_excel_ds_env_override(monkeypatch: pytest.MonkeyPatch) -> None:
    settings = Settings(
        profile="qa",
        base_url="https://example/",
        config_path=Path("/tmp/qa.yaml"),
        data={},
    )
    monkeypatch.setenv("CRDC_EXCEL_DS_VALIDATION", "1")
    assert excel_ds_enabled(settings) is True
